# -*- coding: utf-8 -*-
from docx import Document
PATH = r"d:\python\Financial\docs\参赛报告——房地产产业链风险传导分析金融智能体项目（辛苦学长们进一步补充）_filled.docx"
doc = Document(PATH)
for p in doc.paragraphs:
    if p.text.startswith("供给侧采用标准Ghosh模型"):
        # 找到含该文本的 run 并追加措辞
        for r in p.runs:
            if "初始投入收缩沿产业关联网络" in r.text:
                r.text = r.text.replace(
                    "初始投入收缩沿产业关联网络向下游的传导",
                    "初始投入收缩（供给约束型冲击）沿产业关联网络向下游的传导")
                print("patched")
                break
        break
doc.save(PATH)
# 验证
d2 = Document(PATH)
full = "\n".join(p.text for p in d2.paragraphs)
print("供给约束:", "供给约束型冲击" in full)
print("篡改量化情景参数:", "篡改量化情景参数" in full)
print("请补充残留:", full.count("（请补充）"))