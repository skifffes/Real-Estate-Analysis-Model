# -*- coding: utf-8 -*-
# build docx from report_src/part*.txt
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(r"d:\python\Financial\docs\report_src")
FIG = Path(r"d:\python\Financial\docs\figures")
OUT = Path(r"d:\python\Financial\docs\canpushi_report_v1.docx")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
style.paragraph_format.line_spacing = 1.5

for level, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12.5)]:
    st = doc.styles[level]
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")

def add_p(text, bold=False, center=False, size=None, italic=False):
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)

def add_figure(rel, caption):
    path = FIG / Path(rel).name
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path), width=Cm(15))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.font.size = Pt(10)
    run.bold = True

def render_line(line):
    if line.startswith("#H1 "):
        doc.add_heading(line[4:], level=1)
    elif line.startswith("#H2 "):
        doc.add_heading(line[4:], level=2)
    elif line.startswith("#H3 "):
        doc.add_heading(line[4:], level=3)
    elif line.startswith("#TITLE "):
        add_p(line[7:], bold=True, center=True, size=26)
    elif line.startswith("#TITLE2 "):
        add_p(line[8:], bold=True, center=True, size=22)
    elif line.startswith("#SUB "):
        add_p(line[5:], center=True, size=11, italic=True)
    elif line.startswith("#CENTERBOLD "):
        add_p(line[12:], bold=True, center=True)
    elif line.startswith("#CENTER "):
        add_p(line[8:], center=True, italic=True)
    elif line.startswith("#PBOLD "):
        add_p(line[7:], bold=True)
    elif line.startswith("#BULLET "):
        doc.add_paragraph(line[8:], style="List Bullet")
    elif line.startswith("#FIG "):
        rel, cap = line[5:].split("|", 1)
        add_figure(rel.strip(), cap.strip())
    elif line.startswith("#PAGEBREAK"):
        doc.add_page_break()
    elif line.startswith("#P "):
        add_p(line[3:])
    elif line.strip():
        add_p(line)

for f in sorted(SRC.glob("part*.txt")):
    for line in f.read_text(encoding="utf-8").splitlines():
        if line.strip():
            render_line(line)

doc.save(str(OUT))
print("saved:", OUT, OUT.stat().st_size // 1024, "KB")