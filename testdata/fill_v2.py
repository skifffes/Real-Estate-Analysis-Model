# -*- coding: utf-8 -*-
# 参赛报告填充 v2：在学长模板中插入全部内容（保持原样式）
import copy
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

SRC = r"d:\python\Financial\docs\参赛报告——房地产产业链风险传导分析金融智能体项目（辛苦学长们进一步补充）.docx"
FIGDIR = r"d:\python\Financial\docs\figures"

doc = Document(SRC)

def insert_after(para, text=None, bold=None, center=None, size=None, picture=None):
    """在 para 之后插入新段落（复制其样式），返回新段落对象"""
    new_p = copy.deepcopy(para._p)
    para._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    np = Paragraph(new_p, para._parent)
    # 清空 runs
    for r in list(np.runs):
        r._element.getparent().remove(r._element)
    if text is not None:
        run = np.add_run(text)
        if bold is not None:
            run.bold = bold
        if size is not None:
            run.font.size = Pt(size)
    if center:
        np.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if picture:
        run = np.add_run()
        run.add_picture(picture, width=Cm(15))
        np.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return np


def find_idx(prefix, start=0):
    for i in range(start, len(doc.paragraphs)):
        if doc.paragraphs[i].text.strip().startswith(prefix):
            return i
    return -1

# ============ 1. 参赛项目简介 ============
i = find_idx("参赛项目简介")
fill_p = doc.paragraphs[i + 1]  # （请补充）
INTRO = [
    "房地产通常被认为具有产业链长、关联行业多、对宏观经济影响范围广等特点。钢铁、水泥、建筑施工、机械设备、家居家电、金融服务等行业，都与房地产开发和交易活动存在不同程度的联系。但“关联行业多”并不等于房地产对所有行业都具有同样强的带动作用；房地产市场发生变化后，风险也不会沿着一条固定路径、以同样速度传导到每一个上下游行业。",
    "本项目从这一问题出发，尝试回答三个更具体的问题：第一，房地产在产业链中的真实关联地位如何，其带动能力是否与通常认知一致；第二，当房地产市场出现需求收缩、上游行业出现供给收缩时，风险分别通过什么路径传导；第三，能否把投入产出模型、行业风险指标、历史案例和大语言模型结合起来，将原本较为复杂的产业链分析转化为普通研究者可以直接使用的分析工具。",
    "项目以北京市投入产出表为主要基础数据，将42部门投入产出表按照房地产产业链分析需要聚合为13个部门，并对2012年、2017年和2023年的产业关联结构进行比较。在需求侧，采用Leontief投入产出模型分析最终需求变化向产业链的传导；在供给侧，采用Ghosh模型分析初始投入收缩向下游部门的关联压力。同时，项目结合收入、债务、现金流和市场需求等因素构建情景风险评分，并通过历史案例和月度指标辅助判断风险传导的阶段和路径。",
    "在研究框架基础上，项目进一步开发了“房地产产业链风险分析智能体”。系统并不让大语言模型直接生成计算结果，而是将知识库检索、Leontief模型、Ghosh模型、风险评分和案例检索封装为独立工具，由服务端判断模型适用条件并完成确定性计算，大语言模型主要承担问题理解、工具调度和结果解释工作。当前系统包含智能问答、风险仪表盘、数据上传和报告中心等功能，能够将产业链分析结果以图谱、评分、阶段判断和结构化报告等形式展示。现有知识库包含60余篇资料并形成8700余个检索切片，案例库整理了9个结构化风险案例。",
    "本项目的重点不是利用人工智能替代产业经济分析，而是将经过验证的研究方法封装成一个可交互、可复现、能够主动说明模型边界的金融分析工具。",
]
fill_p.text = INTRO[0]
cur = fill_p
for txt in INTRO[1:]:
    cur = insert_after(cur, text=txt)
print("1. intro filled")

# ============ 2. （四）实施过程：在 2. Ghosh 段后补图1、图2 ============
i = find_idx("在此基础上，项目增加供给侧分析模块")
cur = insert_after(doc.paragraphs[i], picture=FIGDIR + r"\fig1_multiplier.png")
cur = insert_after(cur, text="图1  北京市主要行业影响力乘数演变（2012/2017/2023）", center=True, bold=True, size=10)
cur = insert_after(cur, text="图1显示，2012、2017、2023三个年份中，房地产部门自身的影响力乘数（2.23、1.84、2.23）始终低于建筑业（4.32、4.11、3.84）、钢铁（4.98、4.64、4.75）与建材（3.81、3.84、3.83）。这一“反常识”结果说明：房地产产业链的带动效应并非通过其自身最高的乘数实现，而是作为需求组织者，首先改变建筑业与专业服务需求，再经由这些部门向钢铁、建材、家电家具等方向二次传导。", size=10.5)
cur = insert_after(cur, picture=FIGDIR + r"\fig2_framework.png")
cur = insert_after(cur, text="图2  项目总体研究思路与模型分流框架", center=True, bold=True, size=10)
print("2. fig1/fig2 inserted after Ghosh para")

# ============ 3. 5. 系统测试与演示结果 后补图3 ============
i = find_idx("从目前结果看，项目已经基本完成")
cur = insert_after(doc.paragraphs[i], picture=FIGDIR + r"\fig3_stages.png")
cur = insert_after(cur, text="图3  房地产需求侧风险三阶段传导时间轴", center=True, bold=True, size=10)
print("3. fig3 inserted")

# ============ 4. 五、主要结论及不足（请补充） ============
i = find_idx("五、主要结论及不足")
h = doc.paragraphs[i]
CONCL = [
    "结合多年份投入产出分析、模型测算和历史案例，本项目形成四点主要结论。",
    "第一，房地产产业链“关联广”不等于“房地产自身乘数最高”。2012、2017和2023三年数据均显示，房地产影响力乘数低于建筑业、钢铁和建材。房地产对经济的影响不能只用自身乘数衡量，它更重要的特征是连接建筑、原材料、金融、服务和后周期消费等多个部门，并通过这些节点继续向外传导。这一结果构成本项目最重要的“反常识”发现。",
    "第二，房地产产业链具有稳定性，但关联强度随结构变化调整。三个年份的房地产上游投入方向高度连续，但专业服务直接投入系数由2012年0.109升至2023年0.162，金融投入由0.137降至0.069，建筑、家电等完全关联乘数亦有变化。产业链不是静态网络，需要持续更新。",
    "第三，房地产风险是多阶段、异质化的传导过程。风险可能先表现为供应商回款与现金流问题，再演化为工程停滞、原材料需求下降与后周期消费缺失，最终进入行业出清与结构重构。判断风险需要同时观察应收账款、企业现金流、新开工、原材料产量、竣工及后周期消费等指标。",
    "第四，不同性质的冲击需要不同分析方法。Leontief适合最终需求数量变化，Ghosh用于供给侧初始投入变化，房价与原材料价格变化不直接套入数量框架。项目形成的不是“所有问题都能输出数字”的系统，而是能够区分“什么时候应该算、什么时候不应该算”的分析框架。模型边界本身是结果可靠性的组成部分。",
    "项目不足主要包括：地区投入产出表中间使用含调入与进口因素，Leontief与Ghosh结果目前解释为“产业关联压力情景测算”而非实际产出预测；价格类冲击尚缺专门的投入产出价格模型；情景风险评分中收入、现金流与需求分项为规则映射的代理变量，待接入真实行业财务数据校准；三阶段时间窗口来自案例归纳，尚需月度数据的领先滞后分析验证。",
]
cur = h
for txt in CONCL:
    cur = insert_after(cur, text=txt)
print("4. conclusion filled")

# ============ 5. 六、下一步方向（请补充） ============
i = find_idx("六、下一步方向")
h = doc.paragraphs[i]
NEXT = [
    "一是引入投入产出价格模型与价格-需求弹性模块，将原材料涨价与房价变化纳入量化框架，与现有数量模型形成“数量-价格”双轨分析。",
    "二是提取地区投入产出表中的调入与进口明细，构造区内投入系数，区分本地生产与外部供给，提高Leontief与Ghosh测算的地区解释力。",
    "三是接入上市公司与行业层面的真实营收、经营现金流和应收账款数据，对情景风险评分的权重与阈值进行历史回测，使评分从横向比较指标向产业风险监测指标演进。",
    "四是基于销售、新开工、施工、水泥产量、家电消费等月度序列开展领先滞后分析，实证估计产业链各环节的传导时滞，替代目前基于案例归纳的时间窗口。",
    "五是扩展全国与多区域投入产出表，比较不同地区房地产风险传导结构差异，并研究冲击经由供应链的跨区域传递。",
]
cur = h
for txt in NEXT:
    cur = insert_after(cur, text=txt)
print("5. next-direction filled")

doc.save(SRC.replace(".docx", "_filled.docx"))
print("saved: ..._filled.docx")