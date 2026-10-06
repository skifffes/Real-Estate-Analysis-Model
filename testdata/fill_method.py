# -*- coding: utf-8 -*-
# 补齐：模板第四章末尾插入 GPT 版双模型公式与口径段（学长内容之后）
import copy
from docx import Document
from docx.shared import Pt

PATH = r"d:\python\Financial\docs\参赛报告——房地产产业链风险传导分析金融智能体项目（辛苦学长们进一步补充）_filled.docx"
doc = Document(PATH)
paras = doc.paragraphs

def insert_after(para, text=None, bold=None, italic=None, center=None):
    new_p = copy.deepcopy(para._p)
    para._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    np = Paragraph(new_p, para._parent)
    for r in list(np.runs):
        r._element.getparent().remove(r._element)
    if text is not None:
        run = np.add_run(text)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic
    if center:
        np.alignment = 1
    return np

# 锚点：第四章末（“但这些技术并不是研究的目的”段 / 或“5. 系统测试与演示结果”末段）
anchor_idx = next(i for i, p in enumerate(paras) if p.text.startswith("相比直接向通用大模型提问"))
anchor = paras[anchor_idx]

BLOCK = [
    ("P", "在上述实现基础上，项目对双模型的数学口径与适用边界作进一步规范。需求侧采用Leontief模型：ΔX = (I - A)^(-1) ΔY，其中最终需求冲击情景设定为“房地产—建筑业联合最终需求等幅变化”的复合压力测试——它并非对单一现实指标的一一映射，而是用于比较现有产业结构下各部门相对暴露程度的统一压力情景；相应结果解释为“需求端产业关联压力的情景测算”，而非北京市各行业实际产出的确定性预测。"),
    ("P", "供给侧采用标准Ghosh模型：构造供给分配系数矩阵 B = D^(-1) Z（按卖方部门总产出系数化，B_ij = Z_ij / X_i），通过 ΔX = ΔV (I - B)^(-1) 刻画初始投入收缩沿产业关联网络向下游的传导。由于 B = D^(-1)AD 与直接消耗矩阵A互为相似矩阵，两者谱半径相同（本项目数据下约为0.7303），I-B数值稳定可逆。"),
    ("P", "需要特别说明地区表口径：北京市地区投入产出表的中间使用包含跨地区调入和进口（行平衡关系为中间使用+最终使用-进口=总产出），因此B不宜简单解释为“北京市本地产出的销售分配比例”，而更适合作为各下游部门对相关产品的中间使用暴露与产业关联强度指标；相应地，模型结果解释为供给侧投入压力沿产业链传导的情景测算，而非对实际产出变化的确定性预测。价格与资产价格冲击（原材料涨价、房价变化）属价格效应，不纳入上述数量模型，系统仅输出定性机制分析，精确量化所需的投入产出价格模型与价格-需求弹性模型留作后续扩展。"),
    ("PBOLD", "上述“冲击识别—模型适用边界—情景测算—定性补充”的分流设计，是本系统区别于通用问答系统的关键：所有数量结果均由确定性引擎在服务端计算并强制覆盖大语言模型输出，大语言模型不决定模型适用边界，也不能篡改量化情景参数。"),
]
cur = anchor
for kind, txt in BLOCK:
    cur = insert_after(cur, text=txt, bold=(kind == "PBOLD"))
print("4. methodology block inserted after para", anchor_idx)

doc.save(PATH)
print("saved")