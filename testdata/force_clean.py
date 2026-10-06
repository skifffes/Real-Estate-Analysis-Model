# -*- coding: utf-8 -*-
# 段落级强制清理“（请补充）”后缀（处理跨run拆分）
from docx import Document

PATH = r"d:\python\Financial\docs\参赛报告——房地产产业链风险传导分析金融智能体项目（辛苦学长们进一步补充）_filled.docx"
doc = Document(PATH)
n = 0
for p in doc.paragraphs:
    if "（请补充）" in p.text:
        full_clean = p.text.replace("（请补充）", "")
        # 保留第一个 run 的格式，其余 run 清空
        if p.runs:
            p.runs[0].text = full_clean
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.text = full_clean
        n += 1
        print("cleaned:", p.text[:36])
doc.save(PATH)
print("total cleaned:", n)

# 复验
d2 = Document(PATH)
resid = sum(1 for p in d2.paragraphs if "（请补充）" in p.text)
print("residual:", resid)